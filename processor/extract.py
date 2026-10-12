"""Bounded document extraction. stdout is JSON; diagnostics go to stderr.
No model credentials, database access, or network are used in this process.
"""
import argparse
import base64
import io
import json
import sys
import zipfile
from pathlib import Path

MAX_PAGES = 250
MAX_CHARS = 2_000_000
MAX_IMAGES = 100
MAX_PIXELS = 25_000_000
MAX_OFFICE_EXPANDED_BYTES = 1024 * 1024 * 1024


def image_record(raw, page):
    from PIL import Image, ImageOps
    Image.MAX_IMAGE_PIXELS = MAX_PIXELS
    with Image.open(io.BytesIO(raw)) as image:
        if image.width * image.height > MAX_PIXELS:
            raise ValueError("Image exceeds 25 million pixels")
        image = ImageOps.exif_transpose(image).convert("RGB")
        image.thumbnail((1600, 1600))
        output = io.BytesIO()
        image.save(output, format="JPEG", quality=85)
        return {"text": "", "page": page, "kind": "image", "image": base64.b64encode(output.getvalue()).decode(), "mime": "image/jpeg"}


def extract(path, filename, include_images):
    suffix = Path(filename).suffix.lower()
    records = []
    presentation_has_images = False
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        records.append(image_record(Path(path).read_bytes(), 1))
    elif suffix in {".txt", ".md", ".csv"}:
        raw = Path(path).read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("gb18030")
        records.append({"text": text, "page": 1, "kind": "text"})
    elif suffix == ".pdf":
        import pymupdf as fitz
        with fitz.open(path) as document:
            if document.is_encrypted:
                raise ValueError("Password-protected PDF is not supported")
            if len(document) > MAX_PAGES:
                raise ValueError("PDF exceeds 250 pages; split the document first")
            for index, page in enumerate(document):
                text = page.get_text(sort=True).strip()
                if text:
                    records.append({"text": text, "page": index + 1, "kind": "text"})
                # Render the page: preserves chart labels and relationships that extracted bitmaps lose.
                if include_images and (page.get_images() or len(text) < 80 or page.get_drawings()):
                    if sum(r["kind"] == "image" for r in records) >= MAX_IMAGES:
                        raise ValueError("Document exceeds 100 visual pages; split it first")
                    pix = page.get_pixmap(matrix=fitz.Matrix(min(1.8, 1600 / max(page.rect.width, page.rect.height)), min(1.8, 1600 / max(page.rect.width, page.rect.height))))
                    records.append(image_record(pix.tobytes("png"), index + 1))
                elif not text:
                    raise ValueError("Scanned PDF requires a configured VISION_MODEL")
    elif suffix in {".docx", ".pptx", ".xlsx"}:
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_OFFICE_EXPANDED_BYTES:
                raise ValueError("Office archive expands beyond 1 GiB")
        if suffix == ".docx":
            from docx import Document
            document = Document(path)
            text = "\n".join([p.text for p in document.paragraphs] + ["\n".join(" | ".join(c.text for c in row.cells) for row in t.rows) for t in document.tables])
            records.append({"text": text, "page": 1, "kind": "text"})
        elif suffix == ".pptx":
            from pptx import Presentation
            slides = Presentation(path).slides
            if len(slides) > MAX_PAGES:
                raise ValueError("Presentation exceeds 250 slides")
            for index, slide in enumerate(slides):
                parts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text)
                    if shape.has_table:
                        parts.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
                    if shape.shape_type == 13:
                        presentation_has_images = True
                        if include_images:
                            records.append(image_record(shape.image.blob, index + 1))
                records.append({"text": "\n".join(parts), "page": index + 1, "kind": "text"})
        else:
            from openpyxl import load_workbook
            # Worker/storage paths have no XLSX suffix. Pass bytes via a file stream so
            # openpyxl validates the archive instead of rejecting the temporary filename.
            # Normal mode exposes merged headers needed for row-aware table extraction.
            with Path(path).open("rb") as source:
                book = load_workbook(source, read_only=False, data_only=True)
                try:
                    if len(book.worksheets) > MAX_PAGES:
                        raise ValueError("Spreadsheet exceeds 250 worksheets")
                    from tables import extract_workbook
                    records.extend(extract_workbook(book))
                finally:
                    book.close()
        if suffix == ".docx" and include_images:
            with zipfile.ZipFile(path) as archive:
                for name in archive.namelist():
                    if name.startswith("word/media/") and Path(name).suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}:
                        records.append(image_record(archive.read(name), 1))
    else:
        raise ValueError("Unsupported file type")
    if sum(len(r["text"]) for r in records) > MAX_CHARS or sum(r["kind"] == "image" for r in records) > MAX_IMAGES:
        raise ValueError("Document exceeds extraction limits; split it first")
    if not any(r["text"].strip() or r.get("image") for r in records):
        if presentation_has_images and not include_images:
            raise ValueError("该 PPTX 为图片型演示文稿，请在 RAG 模型配置中启用并保存视觉模型后重新索引（VISION_MODEL）")
        raise ValueError("No extractable content found")
    return records


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("path")
    parser.add_argument("filename")
    parser.add_argument("--images", action="store_true")
    args = parser.parse_args()
    try:
        if sys.platform == "linux":
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (2 * 1024**3, 2 * 1024**3))
            resource.setrlimit(resource.RLIMIT_CPU, (110, 110))
            resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
        print(json.dumps(extract(args.path, args.filename, args.images), ensure_ascii=False))
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        sys.exit(1)
