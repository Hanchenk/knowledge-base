import tempfile
import unittest
import json
import subprocess
import sys
from pathlib import Path

from extract import extract


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)

    def tearDown(self):
        self.directory.cleanup()

    def test_pdf_preserves_page_and_requires_vision_for_scans(self):
        import pymupdf
        from PIL import Image
        document = pymupdf.open()
        document.new_page().insert_text((50, 50), "Travel reimbursement: submit in 15 days.")
        pdf = self.root / "policy.pdf"
        document.save(pdf)
        document.close()
        records = extract(pdf, "policy.pdf", False)
        self.assertIn("15 days", records[0]["text"])
        self.assertEqual(records[0]["page"], 1)
        image = self.root / "scan.png"
        Image.new("RGB", (300, 200), "white").save(image)
        document = pymupdf.open()
        page = document.new_page()
        page.insert_image(page.rect, filename=str(image))
        scan = self.root / "scan.pdf"
        document.save(scan)
        document.close()
        with self.assertRaisesRegex(ValueError, "VISION_MODEL"):
            extract(scan, "scan.pdf", False)
        records = extract(scan, "scan.pdf", True)
        self.assertTrue(records[0]["image"])
        self.assertEqual(records[0]["kind"], "image")

    def test_docx_extracts_tables_and_images(self):
        from docx import Document
        from PIL import Image
        document = Document()
        document.add_paragraph("新人第一周完成安全培训")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "培训预算"
        table.cell(0, 1).text = "3000 元"
        image = self.root / "diagram.png"
        Image.new("RGB", (80, 60), "white").save(image)
        document.add_picture(str(image))
        path = self.root / "guide.docx"
        document.save(path)
        records = extract(path, path.name, True)
        self.assertIn("3000 元", records[0]["text"])
        self.assertTrue(any(r["kind"] == "image" for r in records))

    def test_office_tables_and_slide_numbers(self):
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches
        book = Workbook()
        book.active.title = "预算"
        book.active.append(["住宿", 600])
        path = self.root / "budget.xlsx"
        book.save(path)
        records = extract(path, path.name, False)
        self.assertIn("住宿 | 600", records[0]["text"])
        self.assertEqual(records[0]["kind"], "table")
        presentation = Presentation()
        for index in range(2):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text = f"Milestone {index + 1}"
        path = self.root / "roadmap.pptx"
        presentation.save(path)
        records = extract(path, path.name, False)
        self.assertEqual([r["page"] for r in records], [1, 2])
        self.assertIn("Milestone 2", records[1]["text"])

    def test_unicode_and_empty_content(self):
        path = self.root / "notes.md"
        path.write_text("# 报销流程\n资料需要审核", encoding="utf-8")
        self.assertIn("资料需要审核", extract(path, path.name, False)[0]["text"])
        path.write_text("  ")
        with self.assertRaisesRegex(ValueError, "No extractable"):
            extract(path, path.name, False)

    def test_image_only_pptx_requires_vision_and_preserves_slide_number(self):
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches
        image = self.root / "slide.png"
        Image.new("RGB", (320, 180), "white").save(image)
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(image), Inches(0), Inches(0))
        path = self.root / "image-slides.pptx"
        presentation.save(path)
        with self.assertRaisesRegex(ValueError, "VISION_MODEL"):
            extract(path, path.name, False)
        images = [r for r in extract(path, path.name, True) if r["kind"] == "image"]
        self.assertEqual(len(images), 1)
        self.assertEqual(images[0]["page"], 2)
        self.assertTrue(images[0]["image"])

    def test_xlsx_worker_cli_reads_extensionless_temporary_file(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = "日报"
        book.active.append(["日期", "发电量"])
        book.active.append(["2025-12-01", 1250.5])
        sheet = book.create_sheet("汇总")
        sheet.append(["合计", 1250.5])
        # The Rust worker uses an opaque temporary path and passes the original filename separately.
        path = self.root / "atlas-temporary-upload"
        book.save(path)
        book.close()
        result = subprocess.run(
            [sys.executable, str(Path(__file__).with_name("extract.py")), str(path), "能源日报.xlsx"],
            capture_output=True, text=True, encoding="utf-8", check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        records = json.loads(result.stdout)
        self.assertEqual(set(r["page"] for r in records), {1, 2})
        self.assertTrue(all(r["kind"] == "table" for r in records))
        self.assertIn("工作表：日报", records[0]["text"])
        self.assertTrue(any("日期：2025-12-01" in r["text"] and "发电量：1250.5" in r["text"] for r in records))
        self.assertIn("工作表：汇总", records[-1]["text"])
        self.assertIn("合计 | 1250.5", records[-1]["text"])

    def test_daily_rows_keep_merged_headers_units_and_month_end_snapshot(self):
        from openpyxl import Workbook
        book = Workbook()
        book.remove(book.active)
        for day, cumulative in [(1, 10), (30, 240)]:
            sheet = book.create_sheet(str(day))
            sheet.append(["日报"])
            sheet.append([f"日期：2025/11/{day}"])
            sheet.append(["场站", "日发电量", "月发电量", "月预算完成率", "月损失电量/万度", None])
            sheet.merge_cells("E3:F3")
            sheet.append([None, "万度", "万度", "%", "限电", "故障"])
            sheet.append(["甲站", 10, cumulative, 0.8, 1.5, 2])
            sheet.append(["乙站", 15, cumulative + 30, 1.2, 0, 0])
            sheet["D5"].number_format = "0.00%"
            sheet["D6"].number_format = "0.00%"
            sheet.append(["备注：10月事故"])
        reserved = book.create_sheet("WpsReserved_CellImgList")
        reserved["A200"] = None
        path = self.root / "daily.xlsx"
        book.save(path)
        records = extract(path, path.name, False)
        summary = next(r for r in records if r["kind"] == "table_summary")
        self.assertIn("2025年11月", summary["text"])
        self.assertIn("甲站 | 240 | 80%", summary["text"])
        self.assertIn("乙站 | 270 | 120%", summary["text"])
        self.assertIn("2/30", summary["text"])
        self.assertNotIn("甲站 | 250", summary["text"])
        row = next(r for r in records if "- 场站：甲站" in r["text"])
        self.assertIn("日发电量/万度：10", row["text"])
        self.assertIn("月损失电量/万度/故障：2", row["text"])
        self.assertIn("日期：2025-11-01", row["text"])
        self.assertIn("行：5", row["text"])
        self.assertFalse(any("WpsReserved" in r["text"] for r in records))

    def test_text_only_sheet_and_duplicate_station_do_not_lose_or_invent_data(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(["日期：2025/11/30"])
        sheet.append(["场站", "月发电量"])
        sheet.append(["甲站", 20])
        sheet.append(["甲站", 30])
        sheet = book.create_sheet("说明")
        sheet.append(["项目", "说明"])
        sheet.append(["甲站", "以原表合计为准"])
        path = self.root / "duplicates.xlsx"
        book.save(path)
        records = extract(path, path.name, False)
        self.assertFalse(any(r["kind"] == "table_summary" for r in records))
        self.assertTrue(any("以原表合计为准" in r["text"] for r in records))

    def test_row_date_column_does_not_become_workbook_report_date(self):
        from datetime import date
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(["日期", "场站", "月发电量"])
        sheet.append([date(2025, 11, 1), "甲站", 20])
        sheet.append([date(2025, 12, 1), "甲站", 30])
        path = self.root / "rows.xlsx"
        book.save(path)
        records = extract(path, path.name, False)
        self.assertFalse(any(r["kind"] == "table_summary" for r in records))
        self.assertTrue(any("日期：2025-12-01" in r["text"] for r in records))


if __name__ == "__main__":
    unittest.main()
